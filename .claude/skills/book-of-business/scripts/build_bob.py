#!/usr/bin/env python3
"""
Build a Book of Business workbook from canonical rows (output of normalize.py).

Usage:
  python build_bob.py <canonical.csv> <out.xlsx> [--map serial_map.csv|xlsx]
                      [--today YYYY-MM-DD] [--customer "Ethos"]
                      [--resolve resolutions.json]

Writes <out.xlsx> and <out>_summary.json. Every rule lives here so each run
applies them identically. Anything ambiguous is HELD (not guessed) and listed
in the summary under "questions" for the user to answer. Answers go into a
resolutions.json file and the script is re-run with --resolve.

resolutions.json shape (all keys optional):
  {"include_rows": [row_ids], "exclude_rows": [row_ids],
   "assign_site": {"row_id": "Site Name"}}
"""
import csv, json, re, sys, argparse, datetime as dt
from collections import defaultdict, OrderedDict
from pathlib import Path

# ---------------------------------------------------------------- constants
RELABEL = {"cubex connect": "Interface Fees",
           "cubex connect lite": "Interface Fees",
           "pmp lite": "CubexPMP Lite"}
PARTNER_SHORT = {"mitsubishi hc capital america": "Mitsubishi Capital"}
DIRECT = "cubex direct"
ACRONYMS = {"DTLA", "LA", "NOPA", "HQ", "ER", "ICU", "VCA", "NJ", "NY", "NYC",
            "AEC", "USA", "US", "LLC", "PC", "DVM", "II", "III", "IV"}
SMALL = {"of", "and", "the", "at", "in", "on", "for", "a", "an", "to", "by"}
RENEWAL_TYPES = {"renewal", "renewal_expansion", "renewal upgrade", "renewal expansion"}
HEADERS = ["Account Name", "Asset Name", "Go Live Date", "Contract Start Date",
           "Contract End Date", "Renewal Start Date", "Renewal End Date",
           "Monthly Price", "Serial Number", "Contract #", "Financing Partner"]


# ---------------------------------------------------------------- helpers
def blank(v):
    return v is None or str(v).strip() == "" or str(v).strip().lower() in ("nan", "none", "-")


def s(v):
    return "" if blank(v) else str(v).strip()


def num(v):
    if blank(v):
        return None
    try:
        return float(str(v).replace("$", "").replace(",", ""))
    except ValueError:
        return None


def date(v):
    if blank(v):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return v.date() if isinstance(v, dt.datetime) else v
    t = str(v).strip().split(" ")[0].split("T")[0]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def add_months(d, m):
    m = int(round(m))
    y, mo = divmod(d.month - 1 + m, 12)
    y += d.year
    mo += 1
    import calendar
    return dt.date(y, mo, min(d.day, calendar.monthrange(y, mo)[1]))


def months_between(a, b):
    return (b.year - a.year) * 12 + (b.month - a.month) + (1 if b.day >= a.day - 1 else 0)


def key(name):
    return re.sub(r"\s+", " ", s(name)).lower()


def serial_text(v):
    t = s(v)
    return t[:-2] if re.fullmatch(r"\d+\.0", t) else t


def is_hq(name):
    n = s(name)
    return n.startswith("_") or bool(re.search(r"\bHQ\)?\s*$", n)) or "(HQ)" in n.upper()


def title_case(name, flags):
    words, out, unclear = re.split(r"(\s+|-|/|\(|\))", s(name)), [], []
    first = True
    for w in words:
        if not w or re.fullmatch(r"\s+|-|/|\(|\)", w):
            out.append(w); continue
        core = re.sub(r"[^A-Za-z0-9']", "", w)
        if core.upper() in ACRONYMS:
            out.append(w.upper())
        elif core.isupper() and 2 <= len(core) <= 5 and core.isalpha():
            out.append(w); unclear.append(w)          # unknown all-caps token
        elif re.search(r"[a-z][A-Z]", core):
            out.append(w)                               # McDonald, iPhone
        elif core.lower() in SMALL and not first:
            out.append(w.lower())
        elif "'" in w:
            out.append(w[0].upper() + w[1:].lower())
        else:
            out.append(w[:1].upper() + w[1:].lower())
        first = False
    res = "".join(out)
    if unclear:
        flags.append({"kind": "casing_unclear", "site": res, "tokens": unclear,
                      "detail": "Kept original casing for unrecognized all-caps token(s)."})
    return res


def relabel(asset, product):
    base = s(asset) or s(product)
    return RELABEL.get(base.lower(), base)


def partner(name):
    n = s(name)
    return PARTNER_SHORT.get(n.lower(), n)


# ---------------------------------------------------------------- pricing
def current_price(r, today, flags):
    t1p, t2p = num(r["tier1_price"]), num(r["tier2_price"])
    t1d, t2d = num(r["tier1_dur"]), num(r["tier2_dur"])
    start, end = date(r["commence"]), date(r["end"])
    net = num(r["net_sales_price"])
    if t1d and t1d > 0 and t1p is not None and start:
        if start and end and t2d and t2d > 0:
            term = months_between(start, end)
            if abs((t1d + t2d) - term) > 1:
                flags.append({"kind": "tier_term_conflict", "row_id": r["row_id"],
                              "site": r.get("site_display"), "asset": r["asset"],
                              "detail": f"Tier 1 + Tier 2 = {t1d + t2d:g} months, contract term = {term} months."})
        boundary = add_months(start, t1d)
        if today < boundary:
            return t1p, "tier1"
        if t2p is None:
            flags.append({"kind": "tier2_missing", "row_id": r["row_id"], "asset": r["asset"],
                          "detail": "Past Tier 1 boundary but no Tier 2 Price; used Net Sales Price."})
            return net, "net"
        return t2p, "tier2"
    return net, "net"


# ---------------------------------------------------------------- main build
def build(rows, today, smap, resolve):
    flags, questions, excluded, held = [], [], [], []
    raw_count = len(rows)
    inc_force = set(map(str, resolve.get("include_rows", [])))
    exc_force = set(map(str, resolve.get("exclude_rows", [])))
    assign = {str(k): v for k, v in resolve.get("assign_site", {}).items()}

    for i, r in enumerate(rows):
        r["row_id"] = str(i + 2)               # spreadsheet-style row number in canonical csv
        r["serial"] = serial_text(r.get("serial"))
        r["asset"] = relabel(r.get("asset_name"), r.get("product_name"))
        r["type_l"] = s(r.get("type")).lower()
        r["opp_key"] = s(r.get("opp_id")) or key(r.get("opp_name"))

    def drop(r, reason):
        excluded.append({"row_id": r["row_id"], "reason": reason, "account": s(r["account_name"]),
                         "opp": s(r["opp_name"]), "asset": r["asset"], "serial": r["serial"]})

    live = []
    for r in rows:
        if r["row_id"] in exc_force:
            drop(r, "excluded by user resolution"); continue
        if blank(r.get("account_name")):
            drop(r, "no account name (footer/legal row)"); continue
        cd = date(r.get("close_date"))
        if (r["type_l"] in RENEWAL_TYPES and cd and cd > today and blank(r.get("product_name"))
                and blank(r.get("asset_name")) and num(r.get("net_sales_price")) is None):
            drop(r, "placeholder renewal (future close, no product/asset/price)"); continue
        live.append(r)

    # --- Expansion rows: verify at opportunity level before excluding
    non_exp_serials = {r["serial"] for r in live if r["serial"] and r["type_l"] != "expansion"}
    exp_opps = defaultdict(list)
    for r in live:
        if r["type_l"] == "expansion":
            exp_opps[r["opp_key"]].append(r)
    keep_exp = set()
    for ok, lines in exp_opps.items():
        unique = [l["serial"] for l in lines if l["serial"] and l["serial"] not in non_exp_serials]
        if unique or any(l["row_id"] in inc_force for l in lines):
            keep_exp.add(ok)
            flags.append({"kind": "expansion_kept", "opp": s(lines[0]["opp_name"]),
                          "serials": unique,
                          "detail": "Expansion opportunity carries serial(s) not found on any other row; kept."})
    nxt = []
    for r in live:
        if r["type_l"] == "expansion" and r["opp_key"] not in keep_exp:
            drop(r, "Expansion placeholder/duplicate (no unique serial)")
        else:
            nxt.append(r)
    live = nxt

    # --- Renewal chains: find superseded originals
    orig_ids, orig_names, orig_contracts = set(), set(), set()
    for r in live:
        if r["type_l"] in RENEWAL_TYPES:
            if s(r.get("orig_opp_id")): orig_ids.add(s(r["orig_opp_id"]))
            if s(r.get("orig_opp_name")): orig_names.add(key(r["orig_opp_name"]))
            if s(r.get("originating_contract")): orig_contracts.add(s(r["originating_contract"]))
    renewal_opps_by_contract = defaultdict(set)
    for r in live:
        if r["type_l"] in RENEWAL_TYPES and s(r.get("originating_contract")):
            renewal_opps_by_contract[s(r["originating_contract"])].add(r["opp_key"])

    def superseded(r):
        if s(r.get("opp_id")) and s(r["opp_id"]) in orig_ids: return True
        if key(r.get("opp_name")) in orig_names: return True
        c = s(r.get("contract_no"))
        return bool(c) and c in orig_contracts and r["opp_key"] not in renewal_opps_by_contract[c]

    originals, current = [], []
    for r in live:
        (originals if superseded(r) and r["row_id"] not in inc_force else current).append(r)

    # renewals with no originating link, sharing a serial with another current row -> ask
    by_serial = defaultdict(list)
    for r in current:
        if r["serial"]:
            by_serial[r["serial"]].append(r)

    # --- Site keys
    def site_of(r):
        if r["row_id"] in assign:
            return key(assign[r["row_id"]]), assign[r["row_id"]]
        if r["serial"] and smap and r["serial"] in smap:
            return key(smap[r["serial"]]), smap[r["serial"]]
        return key(r["account_name"]), s(r["account_name"])

    for r in current + originals:
        r["site_key"], r["site_raw"] = site_of(r)

    # --- Serial conflicts among current rows
    drop_ids = set()
    for ser, lines in by_serial.items():
        opps = {l["opp_key"] for l in lines}
        if len(opps) < 2:
            continue
        sites = {l["site_key"] for l in lines}
        if len(sites) > 1:
            # keep the site on the newer renewal row, drop older
            lines.sort(key=lambda l: (l["type_l"] in RENEWAL_TYPES,
                                      date(l["commence"]) or dt.date.min,
                                      date(l["close_date"]) or dt.date.min))
            keep = lines[-1]
            for l in lines[:-1]:
                if l["row_id"] in inc_force: continue
                drop_ids.add(l["row_id"])
                flags.append({"kind": "serial_two_sites", "serial": ser, "kept_site": keep["site_raw"],
                              "dropped_site": l["site_raw"], "dropped_row": l["row_id"],
                              "detail": "Same serial at two sites; kept the newer renewal row."})
        else:
            if any(l["row_id"] in inc_force for l in lines):
                for l in lines:
                    if l["row_id"] not in inc_force: drop_ids.add(l["row_id"])
                continue
            questions.append({"kind": "chain_ambiguous", "serial": ser, "site": lines[0]["site_raw"],
                              "rows": [{"row_id": l["row_id"], "opp": s(l["opp_name"]), "type": s(l["type"]),
                                        "commence": s(l["commence"]), "end": s(l["end"])} for l in lines],
                              "ask": "Which row is the current contract for this serial? Answer with include_rows/exclude_rows."})
            for l in lines:
                drop_ids.add(l["row_id"]); held.append(l)
    kept = []
    for r in current:
        if r["row_id"] in drop_ids:
            if r not in held:
                drop(r, "older row for serial found at another site")
        else:
            kept.append(r)
    for r in originals:
        drop(r, "superseded by renewal")

    # --- HQ fee rows: move to physical site only when pairing is unambiguous
    hw_sites_by_contract = defaultdict(set)
    for r in kept:
        if r["serial"] and s(r.get("contract_no")) and not is_hq(r["site_raw"]):
            hw_sites_by_contract[s(r["contract_no"])].add((r["site_key"], r["site_raw"]))
    for r in kept:
        if r["serial"] or not is_hq(r["site_raw"]) or r["row_id"] in assign:
            continue
        c = s(r.get("contract_no"))
        sites = hw_sites_by_contract.get(c, set()) if c else set()
        if len(sites) == 1:
            (sk, sr), = sites
            r["site_key"], r["site_raw"] = sk, sr
            flags.append({"kind": "hq_fee_moved", "row_id": r["row_id"], "asset": r["asset"], "to_site": sr,
                          "detail": f"HQ fee row paired to site by contract {c}."})
        elif len(sites) > 1:
            questions.append({"kind": "hq_fee_ambiguous", "row_id": r["row_id"], "asset": r["asset"],
                              "contract": c, "candidate_sites": sorted(x[1] for x in sites),
                              "ask": "Which site does this HQ fee row belong to? Answer with assign_site."})
            held.append(r)
    kept = [r for r in kept if r not in held]

    # --- Price and dates for each kept row
    orig_by_serial = {o["serial"]: o for o in originals if o["serial"]}
    orig_by_opp_asset = {(s(o.get("opp_id")) or key(o["opp_name"]), key(o["asset"])): o for o in originals}
    orig_by_name_asset = {(key(o["opp_name"]), key(o["asset"])): o for o in originals}
    for r in kept:
        r["site_display"] = title_case(r["site_raw"], flags) if r["site_raw"] else ""
    # dedupe casing flags
    seen = set(); fl2 = []
    for f in flags:
        k = json.dumps(f, sort_keys=True, default=str)
        if k not in seen: seen.add(k); fl2.append(f)
    flags[:] = fl2

    for r in kept:
        is_ren = r["type_l"] in RENEWAL_TYPES
        if is_ren:
            r["c_start"], r["c_end"] = date(r.get("orig_commence")), date(r.get("orig_end"))
            r["r_start"], r["r_end"] = date(r.get("commence")), date(r.get("end"))
            if not r["c_start"] or not r["c_end"]:
                orig = orig_by_serial.get(r["serial"])
                if orig:
                    r["c_start"] = r["c_start"] or date(orig["commence"])
                    r["c_end"] = r["c_end"] or date(orig["end"])
            price, src = current_price(r, today, flags)
            if r["r_start"] and r["r_start"] > today:
                orig = (orig_by_serial.get(r["serial"]) if r["serial"] else None) \
                    or orig_by_opp_asset.get((s(r.get("orig_opp_id")), key(r["asset"]))) \
                    or orig_by_name_asset.get((key(r.get("orig_opp_name")), key(r["asset"])))
                if orig:
                    price, src = current_price(orig, today, flags)
                    src = "original contract (renewal not started)"
                    flags.append({"kind": "renewal_not_started", "row_id": r["row_id"], "site": r["site_display"],
                                  "asset": r["asset"], "renewal_start": str(r["r_start"]),
                                  "detail": "Showing original contract price. Contract # and Financing Partner are the renewal line's own."})
                else:
                    price, src = None, "missing"
                    flags.append({"kind": "orig_price_missing", "row_id": r["row_id"], "site": r["site_display"],
                                  "asset": r["asset"],
                                  "detail": "Renewal has not started and no original line was found; price left blank."})
        else:
            r["c_start"], r["c_end"] = date(r.get("commence")), date(r.get("end"))
            r["r_start"] = r["r_end"] = None
            price, src = current_price(r, today, flags)
        r["price"], r["price_src"] = price, src
        r["partner"] = partner(r.get("finance_co"))
        r["contract"] = s(r.get("contract_no"))
        r["go"] = date(r.get("go_live")) if r["serial"] else None
        if not is_ren and r["c_end"] and r["c_end"] < today:
            flags.append({"kind": "expired_no_renewal", "row_id": r["row_id"], "site": r["site_display"],
                          "asset": r["asset"], "detail": f"Contract ended {r['c_end']:%-m/%-d/%Y}; no renewal found."})

    # --- Blocks
    blocks = OrderedDict()
    for r in kept:
        blocks.setdefault(r["site_key"], []).append(r)
    out_blocks, excl_sites = [], []
    for sk, lines in blocks.items():
        hw = [l for l in lines if l["serial"]]
        if not hw:
            excl_sites.append({"site": lines[0]["site_display"],
                               "rows": [{"row_id": l["row_id"], "asset": l["asset"],
                                         "price": l["price"]} for l in lines]})
            for l in lines:
                drop(l, "site has no main equipment row")
            continue
        fees = [l for l in lines if not l["serial"]]
        name = hw[0]["site_display"]
        out_blocks.append({"site": name, "hq": is_hq(hw[0]["site_raw"]), "rows": hw + fees})
    for l in held:
        excluded.append({"row_id": l["row_id"], "reason": "HELD pending answer", "account": s(l["account_name"]),
                         "opp": s(l["opp_name"]), "asset": l["asset"], "serial": l["serial"]})
    out_blocks.sort(key=lambda b: (b["hq"], b["site"].lower()))
    return out_blocks, excluded, excl_sites, flags, questions, raw_count


# ---------------------------------------------------------------- workbook
def write_xlsx(blocks, path, flags):
    import openpyxl
    from openpyxl.styles import Font, Border, Side, Alignment
    from openpyxl.comments import Comment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master List"
    F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    side = Side(style="thin", color="BFBFBF")
    box = Border(left=side, right=side, top=side, bottom=side)
    DATE, CUR = "m/d/yyyy", "$#,##0.00"
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, c, h); cell.font = FB; cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    row = 2
    missing = []

    def note(cell, text):
        cell.comment = Comment(text, "BoB Builder")

    for b in blocks:
        first = row
        for l in b["rows"]:
            hwrow = bool(l["serial"])
            vals = [b["site"] if hwrow else None, l["asset"], l["go"] if hwrow else None,
                    l["c_start"], l["c_end"], l["r_start"], l["r_end"], l["price"],
                    l["serial"] if hwrow else None, l["contract"] or None, l["partner"] or None]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row, c, v); cell.font = F; cell.border = box
                if c in (3, 4, 5, 6, 7): cell.number_format = DATE
                if c == 8: cell.number_format = CUR
                if c == 9: cell.number_format = "@"
            if l["price"] is None:
                note(ws.cell(row, 8), "Monthly price missing in Salesforce data. Left blank.")
                missing.append({"kind": "price_missing", "site": b["site"], "asset": l["asset"], "row_id": l["row_id"]})
            if hwrow and l["go"] is None:
                note(ws.cell(row, 3), "Go-live date missing in Salesforce data.")
                missing.append({"kind": "golive_missing", "site": b["site"], "asset": l["asset"], "row_id": l["row_id"]})
            if not l["contract"] and l["partner"] and l["partner"].lower() != DIRECT:
                note(ws.cell(row, 10), "Financed row with no Finance Partner Contract #.")
                missing.append({"kind": "contract_missing", "site": b["site"], "asset": l["asset"], "row_id": l["row_id"]})
            if l["c_start"] is None or l["c_end"] is None:
                note(ws.cell(row, 4), "Contract start/end missing in Salesforce data.")
                missing.append({"kind": "contract_dates_missing", "site": b["site"], "asset": l["asset"], "row_id": l["row_id"]})
            l["xl_row"] = row
            row += 1
        ws.cell(row, 1, "Total").font = FB
        t = ws.cell(row, 8, f"=SUM(H{first}:H{row - 1})")
        t.font = FB; t.number_format = CUR
        for c in range(1, 12):
            ws.cell(row, c).border = Border(left=side, right=side, top=side, bottom=Side(style="medium", color="808080"))
        b["total_row"], b["first_row"] = row, first
        row += 2
    widths = [42, 26, 13, 14, 14, 14, 14, 14, 16, 18, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    wb.save(path)
    return missing


# ---------------------------------------------------------------- validation
def validate(blocks, excluded, raw_count, smap):
    fails = []
    for b in blocks:
        tot = sum(l["price"] or 0 for l in b["rows"])
        ctm = {round(num(l.get("cust_total_monthly")) or -1, 2) for l in b["rows"]}
        b["total"] = round(tot, 2)
        if not any(abs(tot - c) <= 0.01 for c in ctm if c >= 0):
            fails.append({"check": "total_ties_to_customer_total_monthly", "site": b["site"],
                          "block_total": round(tot, 2), "customer_total_monthly_values": sorted(c for c in ctm if c >= 0),
                          "note": "Expected on multi-unit contracts. Flagged, not fixed."})
        hw = b["rows"][0]
        if (hw["c_start"] is None and (s(hw.get("commence")) or s(hw.get("orig_commence")))) or \
           (not hw["partner"] and s(hw.get("finance_co"))):
            fails.append({"check": "dates_or_partner_dropped", "site": b["site"]})
    ser_sites = defaultdict(set)
    for b in blocks:
        for l in b["rows"]:
            if l["serial"]: ser_sites[l["serial"]].add(b["site"])
    for ser, st in ser_sites.items():
        if len(st) > 1:
            fails.append({"check": "duplicate_serial_across_sites", "serial": ser, "sites": sorted(st)})
    if smap:
        for ser in smap:
            if ser in ser_sites and len(ser_sites[ser]) != 1:
                fails.append({"check": "mapped_serial_not_in_one_block", "serial": ser})
    filtered = sum(len(b["rows"]) for b in blocks)
    ok = filtered + len(excluded) == raw_count
    recon = {"raw": raw_count, "included": filtered, "excluded_or_held": len(excluded), "ties": ok}
    if not ok:
        fails.append({"check": "row_count_reconciliation", **recon})
    return fails, recon


def load_map(path):
    if not path:
        return {}
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hdr, data = [str(h or "").lower() for h in rows[0]], rows[1:]
    else:
        with open(p, newline="", encoding="utf-8-sig") as f:
            rd = list(csv.reader(f))
        hdr, data = [h.lower() for h in rd[0]], rd[1:]
    si = next(i for i, h in enumerate(hdr) if "serial" in h)
    ti = next(i for i, h in enumerate(hdr) if i != si and ("site" in h or "account" in h or "location" in h))
    return {serial_text(r[si]): s(r[ti]) for r in data if not blank(r[si]) and not blank(r[ti])}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("canonical"); ap.add_argument("out")
    ap.add_argument("--map"); ap.add_argument("--today"); ap.add_argument("--customer", default="")
    ap.add_argument("--resolve")
    a = ap.parse_args()
    today = dt.date.fromisoformat(a.today) if a.today else dt.date.today()
    with open(a.canonical, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    smap = load_map(a.map)
    resolve = json.load(open(a.resolve)) if a.resolve else {}
    blocks, excluded, excl_sites, flags, questions, raw = build(rows, today, smap, resolve)
    missing = write_xlsx(blocks, a.out, flags)
    fails, recon = validate(blocks, excluded, raw, smap)
    goliv = [l for b in blocks for l in b["rows"] if l["serial"]]
    if goliv and all(l["go"] is None for l in goliv):
        missing = [m for m in missing if m["kind"] != "golive_missing"]
        missing.append({"kind": "golive_missing_account_level", "detail": "Go-live date blank on every hardware row."})
    map_flags = {}
    if smap:
        raw_serials = {serial_text(r.get("serial")) for r in rows if serial_text(r.get("serial"))}
        map_flags = {"serials_missing_from_map": sorted(raw_serials - set(smap)),
                     "map_serials_not_in_data": sorted(set(smap) - raw_serials)}
    reasons = defaultdict(int)
    for e in excluded: reasons[e["reason"]] += 1
    summary = {
        "customer": a.customer, "as_of": today.isoformat(),
        "sites": len(blocks), "physical_sites": sum(not b["hq"] for b in blocks),
        "hq_blocks": sum(b["hq"] for b in blocks),
        "total_mrr": round(sum(b["total"] for b in blocks), 2),
        "row_reconciliation": recon, "excluded_by_reason": dict(reasons),
        "excluded_sites_no_hardware": excl_sites,
        "questions": questions, "validation_failures": fails,
        "missing_data": missing, "flags": flags, "serial_map": map_flags,
        "excluded_rows": excluded,
    }
    sp = str(Path(a.out).with_suffix("")) + "_summary.json"
    json.dump(summary, open(sp, "w"), indent=2, default=str)
    print(json.dumps({k: summary[k] for k in
                      ("sites", "physical_sites", "hq_blocks", "total_mrr", "row_reconciliation", "excluded_by_reason")},
                     indent=2, default=str))
    print(f"questions={len(questions)} validation_failures={len(fails)} missing={len(missing)} flags={len(flags)}")
    print(f"summary -> {sp}")


if __name__ == "__main__":
    main()
