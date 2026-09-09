import pickle, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

d = pickle.load(open('report_rows.pkl','rb'))

wb = Workbook()
ARIAL = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
BASE = Font(name=ARIAL, size=10)
WARN_FILL = PatternFill('solid', fgColor='FFF2CC')
BAD_FILL = PatternFill('solid', fgColor='FCE4E4')

def write_sheet(ws, rows, widths):
    if not rows: return
    cols = list(rows[0].keys())
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical='center')
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            v = r[c]
            if isinstance(v, datetime.datetime): v = v.date()
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BASE
            if isinstance(v, datetime.date): cell.number_format = 'mm/dd/yyyy'
        act = r.get('Action Needed') or ''
        if act:
            fill = BAD_FILL if 'No SFDC account found' in act else WARN_FILL
            ws.cell(row=i, column=cols.index('Action Needed')+1).fill = fill
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 14)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(rows)+1}'

W = {'Member Name': 38, 'SFDC Account Name': 38, 'City': 16, 'Action Needed': 46,
     'Note': 34, 'SFDC Account ID': 20, 'SFDC Type': 12, 'SFDC State': 14,
     'Match Method': 12, 'Confidence': 12, 'SFDC PSI Join Date': 14,
     'SFDC PSI Termination Date': 16, 'Active Date': 12, 'Cancel Date': 12,
     'On This Month Cancelled Sheet': 16, 'SFDC PSI ID': 10, 'PSI ID': 8}

ws = wb.active; ws.title = 'Summary'
write_sheet(wb.create_sheet('Active Members'), d['active'], W)
write_sheet(wb.create_sheet('New Members'), d['new'], W)
write_sheet(wb.create_sheet('Cancelled Members'), d['cancelled'], W)
write_sheet(wb.create_sheet('Stale PSI in SFDC'), d['stale'], W)

# ---- Summary ----
title_font = Font(name=ARIAL, bold=True, size=14)
h2 = Font(name=ARIAL, bold=True, size=11)
ws['A1'] = 'PSIvet Member List vs Salesforce Account Match'; ws['A1'].font = title_font
ws['A2'] = 'Source: August 2026 EOM Membership List (thru 8/31/2026). Matched Sep 9, 2026.'; ws['A2'].font = BASE
rows = [
    ('', ''),
    ('Sheet', 'What it shows'),
    ('Active Members', '4,900 rows from the full membership list, each matched to a Salesforce account where possible'),
    ('New Members', '27 members added in August 2026 (also present in Active Members)'),
    ('Cancelled Members', '27 members cancelled in August 2026'),
    ('Stale PSI in SFDC', 'Salesforce accounts holding an active PSI ID that is NOT on the current member list'),
    ('', ''),
    ('Match results (Active Members)', ''),
    ('Matched by PSI ID (exact)', "=COUNTIF('Active Members'!N2:N4901,\"PSI ID\")"),
    ('Matched by Phone', "=COUNTIF('Active Members'!N2:N4901,\"Phone\")"),
    ('Matched by Name+Geo', "=COUNTIF('Active Members'!N2:N4901,\"Name+Geo\")"),
    ('No match found', "=COUNTIF('Active Members'!N2:N4901,\"No match\")"),
    ('Total member rows', "=COUNTA('Active Members'!A2:A4901)"),
    ('Rows needing a Salesforce update', "=COUNTIF('Active Members'!Q2:Q4901,\"?*\")"),
    ('', ''),
    ('Match results (New Members)', ''),
    ('Matched (any method)', "=COUNTIF('New Members'!N2:N28,\"PSI ID\")+COUNTIF('New Members'!N2:N28,\"Phone\")+COUNTIF('New Members'!N2:N28,\"Name+Geo\")"),
    ('No match found', "=COUNTIF('New Members'!N2:N28,\"No match\")"),
    ('', ''),
    ('Match results (Cancelled Members)', ''),
    ('Matched by PSI ID (exact)', "=COUNTIF('Cancelled Members'!O2:O28,\"PSI ID\")"),
    ('Already terminated in SFDC', "=COUNT('Cancelled Members'!N2:N28)"),
    ('Need PSI_Termination_Date__c set', "=COUNTIF('Cancelled Members'!R2:R28,\"*Termination*\")"),
    ('', ''),
    ('Reverse check', ''),
    ('SFDC accounts with active PSI ID not on current list', "=COUNTA('Stale PSI in SFDC'!A2:A313)"),
    ('...of which cancelled this month', "=COUNTIF('Stale PSI in SFDC'!G2:G313,\"Yes\")"),
    ('...older discrepancies to review', "=COUNTIF('Stale PSI in SFDC'!G2:G313,\"No\")"),
    ('', ''),
    ('How matching worked', ''),
    ('1. PSI ID', 'Member PSI ID = Account PSI_Unique_ID__c (exact)'),
    ('2. Phone', 'Normalized 10-digit phone match, best candidate by name/zip/city/state agreement'),
    ('3. Name+Geo', 'Name token similarity plus zip/city/state confirmation'),
    ('Confidence', 'exact = ID match; high/medium = review recommended for medium; blank = no match'),
    ('Action Needed column', 'Field-level differences to apply in Salesforce (PSI ID, join date, termination date)'),
]
for i, (a, b) in enumerate(rows, 3):
    ws.cell(row=i, column=1, value=a).font = h2 if (a and not b) or a=='Sheet' else BASE
    if isinstance(b, str) and b.startswith('='):
        c = ws.cell(row=i, column=2); c.value = b; c.font = Font(name=ARIAL, size=10, bold=True)
    else:
        ws.cell(row=i, column=2, value=b).font = BASE
ws.column_dimensions['A'].width = 48
ws.column_dimensions['B'].width = 95

out = 'PSI_Member_SFDC_Match_Aug2026.xlsx'
wb.save(out)
print('saved', out)

# verify column letters used in formulas
import openpyxl
wb2 = openpyxl.load_workbook(out)
for s in ('Active Members','New Members','Cancelled Members','Stale PSI in SFDC'):
    hdr = [c.value for c in wb2[s][1]]
    print(s, {get_column_letter(i+1): h for i, h in enumerate(hdr)})
